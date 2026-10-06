from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor


OUT = Path("супервизия/анонс_воркшопа_семифокусная_модель.docx")


def set_run_style(run, size=11, bold=False, color="232F3E"):
    run.font.name = "DejaVu Sans"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def add_paragraph(document, text, size=11, bold=False, color="232F3E", spacing_after=8):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(spacing_after)
    run = paragraph.add_run(text)
    set_run_style(run, size=size, bold=bold, color=color)
    return paragraph


def add_heading(document, text, level=1):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(10 if level == 1 else 6)
    paragraph.paragraph_format.space_after = Pt(6)
    run = paragraph.add_run(text)
    set_run_style(run, size=16 if level == 1 else 13, bold=True, color="27598B")
    return paragraph


def add_bullets(document, items):
    for item in items:
        paragraph = document.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.space_after = Pt(4)
        run = paragraph.add_run(item)
        set_run_style(run)


def add_numbered(document, items):
    for item in items:
        paragraph = document.add_paragraph(style="List Number")
        paragraph.paragraph_format.space_after = Pt(4)
        run = paragraph.add_run(item)
        set_run_style(run)


def build():
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(8)
    run = title.add_run("Воркшоп «Семифокусная модель супервизии»")
    set_run_style(run, size=20, bold=True, color="27598B")

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(18)
    run = subtitle.add_run("3 часа практики для психологов и психотерапевтов")
    set_run_style(run, size=14, bold=True, color="C69538")

    add_paragraph(
        document,
        "Когда случай «не движется», терапевт устает, клиент вызывает сильные чувства, "
        "а в работе появляются вопросы границ, смысла и ответственности, важно иметь ясную карту, "
        "которая помогает не потеряться в процессе. Семифокусная модель супервизии дает такую карту.",
        size=11.5,
    )
    add_paragraph(
        document,
        "Приглашаем на 3-часовой воркшоп для практикующих психологов и психотерапевтов, "
        "где мы разберем, как смотреть на клиентский случай объемно: видеть клиента, терапевта, "
        "их отношения, супервизионный процесс, состояние специалиста и духовно-нравственное измерение помощи.",
        size=11.5,
    )

    add_heading(document, "Что такое семифокусная модель")
    add_paragraph(
        document,
        "Семифокусная модель супервизии — это расширенная карта профессионального внимания. "
        "Она опирается на классическую шестифокусную модель, где анализируется не только клиент, "
        "но и терапевт, их контакт, реакции супервизируемого, отношения с супервизором и состояние самого супервизора.",
    )
    add_paragraph(
        document,
        "Седьмой фокус добавляет то, что особенно важно для христиански ориентированной практики: "
        "отношения с Богом, свободу клиента, смирение специалиста, различение духовного и психологического, "
        "а также внимательное удержание профессиональных границ без ухода в «спасательство».",
    )
    add_paragraph(
        document,
        "Эта модель помогает разбирать сложные случаи глубже и спокойнее: не сводить проблему только к технике, "
        "не объяснять все духовными причинами и не оставаться один на один с тяжелыми переживаниями в работе.",
        bold=True,
        color="C69538",
    )

    add_heading(document, "Для кого")
    add_bullets(
        document,
        [
            "практикующим психологам и психотерапевтам;",
            "специалистам, работающим с христианскими клиентами;",
            "психологам, которым важно соединять профессиональную позицию, этику и духовное измерение помощи;",
            "участникам супервизионных и интервизионных групп;",
            "специалистам, которые сталкиваются с выгоранием, контрпереносом, сложными случаями и вопросами границ;",
            "тем, кто хочет сделать супервизию более структурной, живой и практически применимой.",
        ],
    )

    add_heading(document, "Что будем разбирать на воркшопе")
    add_paragraph(document, "За 3 часа участники освоят семь фокусов супервизионного внимания:")
    add_numbered(
        document,
        [
            "Клиент: запрос, страдание, ресурсы и риски.",
            "Терапевт: гипотезы, стратегия, методы и профессиональная позиция.",
            "Отношения клиента и терапевта: процесс, перенос, контакт, динамика встречи.",
            "Состояние супервизируемого: чувства, телесные реакции, контрперенос, усталость.",
            "Отношения терапевта и супервизора: что проявляется в супервизии из терапевтического процесса.",
            "Состояние супервизора: реакции, гипотезы, ограничения и ответственность.",
            "Отношения с Богом: молитвенное внимание, свобода клиента, смирение специалиста, различение духовного и психологического.",
        ],
    )

    add_heading(document, "Почему стоит прийти")
    add_bullets(
        document,
        [
            "вы получите понятную схему, по которой можно разбирать реальные случаи уже после встречи;",
            "увидите, где в работе теряется контакт, границы или профессиональная позиция;",
            "научитесь замечать, что принадлежит клиенту, что терапевту, а что проявляется между ними;",
            "сможете бережнее относиться к себе как к специалисту и снижать риск выгорания;",
            "получите язык для обсуждения духовных вопросов без давления на клиента и без подмены психотерапии духовными объяснениями.",
        ],
    )

    add_heading(document, "Формат")
    add_paragraph(
        document,
        "Это не просто лекция. Воркшоп сочетает ясное объяснение модели, практические вопросы "
        "для самодиагностики, работу с примерами из практики и обсуждение типичных ошибок супервизионного анализа.",
    )
    add_paragraph(document, "План встречи:", bold=True)
    add_bullets(
        document,
        [
            "вводная часть: зачем психологу супервизионная карта внимания;",
            "обзор шестифокусной модели;",
            "добавление седьмого фокуса в православном и христианском контексте;",
            "практический разбор клиентского случая по семи фокусам;",
            "обсуждение рисков: спасательство, духовные объяснения вместо анализа, нарушение границ;",
            "итоговая карта вопросов для дальнейшей супервизионной работы.",
        ],
    )

    add_heading(document, "Что участники получат")
    add_bullets(
        document,
        [
            "смогут структурировать сложный клиентский случай по семи фокусам;",
            "будут лучше различать профессиональные, личные и духовные аспекты работы;",
            "научатся замечать контрперенос, усталость и риск выгорания;",
            "смогут удерживать профессиональные границы без потери христианского взгляда на человека;",
            "получат семифокусную карту как инструмент для супервизии, интервизии и саморефлексии.",
        ],
    )

    add_heading(document, "Главный результат")
    add_paragraph(
        document,
        "После воркшопа у вас будет не набор общих идей, а рабочая супервизионная рамка: "
        "как задавать вопросы, куда смотреть в сложном случае, как замечать свои реакции и как сохранять "
        "профессиональную ясность в контакте с клиентом.",
        bold=True,
        color="27598B",
    )

    add_heading(document, "Продолжительность")
    add_paragraph(document, "3 часа.", bold=True, color="C69538")

    add_heading(document, "Ведущий")
    add_paragraph(document, "[Имя ведущего, регалии, профессиональный опыт]", color="5B6876")

    add_heading(document, "Дата и место")
    add_paragraph(document, "[Дата, время, город/онлайн-формат]", color="5B6876")

    add_heading(document, "Участие")
    add_paragraph(document, "[Стоимость, условия регистрации, ссылка или контакт]", color="5B6876")
    add_paragraph(
        document,
        "Запишитесь на воркшоп, если хотите разбирать клиентские случаи глубже, увереннее и бережнее — "
        "с опорой на профессиональную супервизионную модель и христианское понимание человека.",
        size=11,
        bold=True,
        color="C69538",
    )
    add_paragraph(
        document,
        "Количество мест может быть ограничено, чтобы сохранить возможность живого обсуждения и практической работы.",
        size=10.5,
        color="5B6876",
    )

    document.save(OUT)
    return OUT


if __name__ == "__main__":
    path = build()
    print(f"Created {path}")
