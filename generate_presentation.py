from math import cos, pi, sin
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


OUT = Path("supervision_orthodox_psychologists.pptx")

WHITE = RGBColor(255, 255, 255)
INK = RGBColor(35, 47, 62)
MUTED = RGBColor(91, 104, 118)
BLUE = RGBColor(39, 89, 139)
BLUE_LIGHT = RGBColor(226, 239, 250)
GOLD = RGBColor(198, 149, 56)
GOLD_LIGHT = RGBColor(253, 245, 226)
GREEN = RGBColor(73, 135, 104)
GREEN_LIGHT = RGBColor(230, 244, 236)
RED = RGBColor(160, 82, 82)
RED_LIGHT = RGBColor(250, 235, 232)
GRAY_LINE = RGBColor(212, 218, 224)


def setup_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = WHITE
    return slide


def add_text(slide, text, left, top, width, height, size=24, color=INK, bold=False, align=None):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.name = "DejaVu Sans"
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = color
    if align is not None:
        p.alignment = align
    return box


def add_title(slide, title, subtitle=None):
    add_text(slide, title, 0.65, 0.35, 12.0, 0.55, size=25, color=BLUE, bold=True)
    if subtitle:
        add_text(slide, subtitle, 0.67, 0.92, 11.7, 0.38, size=11.5, color=MUTED)
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.65), Inches(1.33), Inches(2.15), Inches(0.03))
    line.fill.solid()
    line.fill.fore_color.rgb = GOLD
    line.line.fill.background()


def add_footer(slide, num):
    add_text(slide, f"{num:02d}", 12.38, 7.05, 0.45, 0.22, size=8.5, color=MUTED, align=PP_ALIGN.RIGHT)


def add_card(slide, left, top, width, height, title, body, fill=BLUE_LIGHT, accent=BLUE, title_size=15, body_size=11):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = RGBColor(235, 239, 244)
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(left), Inches(top), Inches(0.08), Inches(height))
    bar.fill.solid()
    bar.fill.fore_color.rgb = accent
    bar.line.fill.background()
    add_text(slide, title, left + 0.24, top + 0.15, width - 0.42, 0.34, size=title_size, color=accent, bold=True)
    add_text(slide, body, left + 0.24, top + 0.6, width - 0.42, height - 0.75, size=body_size, color=INK)
    return shape


def add_bullets(slide, items, left, top, width, height, size=14, color=INK):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = item
        p.level = 0
        p.font.name = "DejaVu Sans"
        p.font.size = Pt(size)
        p.font.color.rgb = color
        p.space_after = Pt(8)
        p.text = "• " + item
    return box


def add_quote(slide, text, source, left, top, width, height):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
    shape.fill.solid()
    shape.fill.fore_color.rgb = GOLD_LIGHT
    shape.line.color.rgb = RGBColor(244, 229, 192)
    add_text(slide, f"«{text}»", left + 0.25, top + 0.22, width - 0.5, height - 0.62, size=15, color=INK)
    add_text(slide, source, left + 0.25, top + height - 0.38, width - 0.5, 0.22, size=9.5, color=GOLD, bold=True, align=PP_ALIGN.RIGHT)


def add_person_icon(slide, cx, cy, scale, color):
    head = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(cx - 0.12 * scale), Inches(cy - 0.38 * scale), Inches(0.24 * scale), Inches(0.24 * scale))
    head.fill.solid()
    head.fill.fore_color.rgb = color
    head.line.fill.background()
    body = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(cx), Inches(cy - 0.14 * scale), Inches(cx), Inches(cy + 0.2 * scale))
    body.line.color.rgb = color
    body.line.width = Pt(2)
    arms = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(cx - 0.22 * scale), Inches(cy + 0.02 * scale), Inches(cx + 0.22 * scale), Inches(cy + 0.02 * scale))
    arms.line.color.rgb = color
    arms.line.width = Pt(2)
    left_leg = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(cx), Inches(cy + 0.2 * scale), Inches(cx - 0.2 * scale), Inches(cy + 0.48 * scale))
    right_leg = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(cx), Inches(cy + 0.2 * scale), Inches(cx + 0.2 * scale), Inches(cy + 0.48 * scale))
    for leg in (left_leg, right_leg):
        leg.line.color.rgb = color
        leg.line.width = Pt(2)


def add_cross_icon(slide, cx, cy, scale, color):
    v = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(cx), Inches(cy - 0.42 * scale), Inches(cx), Inches(cy + 0.42 * scale))
    h = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(cx - 0.27 * scale), Inches(cy - 0.12 * scale), Inches(cx + 0.27 * scale), Inches(cy - 0.12 * scale))
    for line in (v, h):
        line.line.color.rgb = color
        line.line.width = Pt(3)


def add_focus_node(slide, x, y, n, title, color=BLUE, fill=BLUE_LIGHT):
    node = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x - 1.05), Inches(y - 0.42), Inches(2.1), Inches(0.84))
    node.fill.solid()
    node.fill.fore_color.rgb = fill
    node.line.color.rgb = RGBColor(224, 231, 239)
    circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x - 1.0), Inches(y - 0.29), Inches(0.32), Inches(0.32))
    circle.fill.solid()
    circle.fill.fore_color.rgb = color
    circle.line.fill.background()
    add_text(slide, str(n), x - 0.94, y - 0.245, 0.2, 0.18, size=8, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, title, x - 0.58, y - 0.24, 1.42, 0.43, size=9.2, color=INK, bold=True)


def add_pentagram(slide, cx, cy, radius):
    labels = [
        ("Духовное", GOLD, GOLD_LIGHT),
        ("Когнитивное", BLUE, BLUE_LIGHT),
        ("Социальное", GREEN, GREEN_LIGHT),
        ("Телесное", RED, RED_LIGHT),
        ("Эмоциональное", RGBColor(125, 86, 170), RGBColor(241, 234, 250)),
    ]
    points = []
    for i in range(5):
        angle = -pi / 2 + i * 2 * pi / 5
        points.append((cx + radius * cos(angle), cy + radius * sin(angle)))

    star_order = [0, 2, 4, 1, 3, 0]
    for a, b in zip(star_order, star_order[1:]):
        x1, y1 = points[a]
        x2, y2 = points[b]
        line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
        line.line.color.rgb = GRAY_LINE
        line.line.width = Pt(1.7)

    for (label, color, fill), (x, y) in zip(labels, points):
        halo = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x - 0.47), Inches(y - 0.47), Inches(0.94), Inches(0.94))
        halo.fill.solid()
        halo.fill.fore_color.rgb = fill
        halo.line.color.rgb = RGBColor(230, 234, 238)
        add_text(slide, label, x - 0.62, y - 0.15, 1.24, 0.3, size=8.6, color=color, bold=True, align=PP_ALIGN.CENTER)


def build():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    slides = []

    s = setup_slide(prs)
    slides.append(s)
    add_text(s, "Семифокусная модель супервизии", 0.8, 1.15, 11.8, 0.65, size=31, color=BLUE, bold=True, align=PP_ALIGN.CENTER)
    add_text(s, "для православных психологов", 0.8, 1.95, 11.8, 0.45, size=22, color=GOLD, bold=True, align=PP_ALIGN.CENTER)
    add_text(s, "На основе шестифокусной модели супервизии и пентаграммы Гингера", 1.45, 2.62, 10.4, 0.46, size=15, color=MUTED, align=PP_ALIGN.CENTER)
    add_cross_icon(s, 6.65, 4.05, 1.1, GOLD)
    add_text(s, "Профессиональная поддержка, духовное внимание и безопасность клиента", 2.1, 5.45, 9.1, 0.44, size=16, color=INK, align=PP_ALIGN.CENTER)
    add_footer(s, 1)

    s = setup_slide(prs)
    slides.append(s)
    add_title(s, "Основная идея", "Супервизия как пространство профессионального различения и духовной трезвости")
    add_card(s, 0.8, 1.75, 3.85, 3.85, "Профессиональное", "Разбор практики, гипотез, этики, границ, качества помощи и безопасности клиента.", BLUE_LIGHT, BLUE)
    add_card(s, 4.75, 1.75, 3.85, 3.85, "Личностное", "Осознавание реакций терапевта, контрпереноса, уязвимостей, ресурсов и риска выгорания.", GREEN_LIGHT, GREEN)
    add_card(s, 8.7, 1.75, 3.85, 3.85, "Духовное", "Христоцентрирование работы, молитвенная поддержка и различение без подмены терапии духовничеством.", GOLD_LIGHT, GOLD)
    add_footer(s, 2)

    s = setup_slide(prs)
    slides.append(s)
    add_title(s, "Духовное окормление", "Пастырское руководство не заменяет супервизию, но задает духовный контекст")
    add_card(
        s,
        0.85,
        1.72,
        11.65,
        3.05,
        "Определение",
        "Духовное окормление — это пастырское руководство, наставничество и забота священника (духовника) о спасении души верующего психолога. Оно включает молитвенную поддержку, советы в духовной жизни и регулярные встречи для разрешения внутренних вопросов.",
        GOLD_LIGHT,
        GOLD,
        title_size=17,
        body_size=17,
    )
    add_bullets(
        s,
        [
            "Духовник помогает видеть труд не только как технику, но и как служение.",
            "Психолог сохраняет профессиональные границы и ответственность перед клиентом.",
            "Соборность не отменяет компетентности, этики и клинического мышления.",
        ],
        1.15,
        5.05,
        11.1,
        1.55,
        size=13.5,
    )
    add_footer(s, 3)

    s = setup_slide(prs)
    slides.append(s)
    add_title(s, "Что такое супервизия", "Метод профессиональной поддержки и контроля качества")
    add_card(
        s,
        0.85,
        1.65,
        5.75,
        4.55,
        "Супервизия",
        "Практикующий специалист разбирает свою работу с более опытным коллегой — супервизором. Это помогает повышать качество работы, анализировать сложные случаи, избегать выгорания и обеспечивать безопасность клиента.",
        BLUE_LIGHT,
        BLUE,
        body_size=15.5,
    )
    add_card(
        s,
        6.85,
        1.65,
        5.65,
        4.55,
        "Для православного психолога",
        "К профессиональному анализу добавляется внимание к духовному устроению, смирению, молитве, свободе клиента и недопустимости «спасательства» вместо ответственной помощи.",
        GOLD_LIGHT,
        GOLD,
        body_size=15.5,
    )
    add_footer(s, 4)

    s = setup_slide(prs)
    slides.append(s)
    add_title(s, "Евангельские опоры", "Три акцента: соборность, трезвость, просвещение Духом Святым")
    add_quote(s, "ибо, где двое или трое собраны во имя Мое, там Я посреди них", "Мф. 18:20", 0.8, 1.7, 3.8, 3.5)
    add_quote(s, "если слепой ведет слепого, то оба упадут в яму", "Мф. 15:14", 4.77, 1.7, 3.8, 3.5)
    add_quote(s, "Утешитель же, Дух Святый... научит вас всему", "Ин. 14:26", 8.74, 1.7, 3.8, 3.5)
    add_footer(s, 5)

    s = setup_slide(prs)
    slides.append(s)
    add_title(s, "Личность терапевта как инструмент", "Супервизия помогает видеть, чем именно психолог работает в контакте")
    add_quote(s, "Хороший терапевт не пользуется техниками, он применяет себя в и к ситуации...", "Лора Перлз", 0.9, 1.55, 5.7, 2.0)
    add_quote(s, "Главное воздействие терапевта заключается не столько в его словах, сколько в его личности.", "Ролло Мэй", 6.75, 1.55, 5.7, 2.0)
    add_card(s, 2.3, 4.15, 8.75, 1.55, "Ключевой тезис", "Психолог работает «собой»: своим вниманием, телесностью, эмоциональной зрелостью, мышлением, отношениями и духовной ориентацией.", GREEN_LIGHT, GREEN, body_size=14)
    add_footer(s, 6)

    s = setup_slide(prs)
    slides.append(s)
    add_title(s, "Цель супервизии", "Укрепление специалиста и повышение качества помощи")
    add_card(
        s,
        0.9,
        1.68,
        11.55,
        2.25,
        "Цель",
        "Укрепление на пути духовного и личностного роста, повышение качества профессиональной деятельности, безопасности клиента и развитие компетенций психолога через анализ практики с опытными коллегами и священником.",
        BLUE_LIGHT,
        BLUE,
        title_size=17,
        body_size=16,
    )
    add_bullets(
        s,
        [
            "молитвенная поддержка и духовная трезвость;",
            "обучение и развитие профессионального мышления;",
            "профилактика выгорания и профессиональной деформации;",
            "контроль этических норм и границ ответственности.",
        ],
        1.15,
        4.3,
        11,
        2.0,
        size=14.5,
    )
    add_footer(s, 7)

    s = setup_slide(prs)
    slides.append(s)
    add_title(s, "Задачи супервизии", "Семь направлений работы")
    tasks = [
        ("Духовная поддержка", "Христоцентрирование и расширение понимания Промысла Божьего."),
        ("Профессиональная поддержка", "Разбор случая, поиск подходов и поддержка специалиста."),
        ("Обучение и развитие", "Православная антропология, психология, навыки и стиль."),
        ("Контроль качества", "Этика, эффективность и безопасность помощи."),
        ("Профилактика выгорания", "Контрперенос, истощение, профессиональная деформация."),
        ("Развитие осознанности", "Собственные процессы, реакции и перенос терапевта."),
        ("Безопасность", "Позиция специалиста без «спасательства»."),
    ]
    positions = [(0.8, 1.55), (4.55, 1.55), (8.3, 1.55), (0.8, 3.35), (4.55, 3.35), (8.3, 3.35), (2.7, 5.15)]
    colors = [(GOLD_LIGHT, GOLD), (BLUE_LIGHT, BLUE), (GREEN_LIGHT, GREEN), (RED_LIGHT, RED), (BLUE_LIGHT, BLUE), (GREEN_LIGHT, GREEN), (GOLD_LIGHT, GOLD)]
    for (title, body), (x, y), (fill, accent) in zip(tasks, positions, colors):
        add_card(s, x, y, 3.25, 1.35, title, body, fill, accent, title_size=11.2, body_size=8.8)
    add_footer(s, 8)

    s = setup_slide(prs)
    slides.append(s)
    add_title(s, "Варианты, уровни и формы", "Организация супервизионной работы")
    add_card(s, 0.85, 1.55, 3.75, 4.8, "Варианты", "• соборная молитва\n• индивидуальная\n• групповая\n• балинтовские группы\n• семейная", BLUE_LIGHT, BLUE, body_size=15)
    add_card(s, 4.8, 1.55, 3.75, 4.8, "Уровни", "• для новоначальных христиан\n• для молодых психологов\n• для практикующих коллег\n• принятие мнения другого со смирением и мудростью", GOLD_LIGHT, GOLD, body_size=14.2)
    add_card(s, 8.75, 1.55, 3.75, 4.8, "Формы", "• очная: в процессе сессии\n• заочная: обсуждение случая в кругу коллег\n• очно-заочная: частичное включение и последующий разбор", GREEN_LIGHT, GREEN, body_size=14.2)
    add_footer(s, 9)

    s = setup_slide(prs)
    slides.append(s)
    add_title(s, "Шестифокусная модель супервизии", "Классическая карта внимания супервизора")
    center = slide_center = (6.65, 3.75)
    cshape = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(center[0] - 0.9), Inches(center[1] - 0.55), Inches(1.8), Inches(1.1))
    cshape.fill.solid()
    cshape.fill.fore_color.rgb = WHITE
    cshape.line.color.rgb = GOLD
    cshape.line.width = Pt(2.2)
    add_text(s, "случай\nклиента", center[0] - 0.55, center[1] - 0.24, 1.1, 0.48, size=12.5, color=GOLD, bold=True, align=PP_ALIGN.CENTER)
    focus_labels = [
        "клиент",
        "терапевт",
        "отношения\nклиент-терапевт",
        "состояние\nсупервизируемого",
        "отношения\nтерапевт-супервизор",
        "состояние\nсупервизора",
    ]
    focus_points = [(6.65, 1.8), (9.75, 2.65), (9.75, 4.95), (6.65, 5.85), (3.55, 4.95), (3.55, 2.65)]
    for i, ((x, y), label) in enumerate(zip(focus_points, focus_labels), start=1):
        add_focus_node(s, x, y, i, label)
        line = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(center[0]), Inches(center[1]), Inches(x), Inches(y))
        line.line.color.rgb = GRAY_LINE
        line.line.width = Pt(1)
    add_footer(s, 10)

    s = setup_slide(prs)
    slides.append(s)
    add_title(s, "Седьмой фокус: отношения с Богом", "Предложение для православной супервизии")
    add_cross_icon(s, 2.0, 2.1, 0.85, GOLD)
    add_text(s, "7", 1.77, 2.88, 0.46, 0.3, size=18, color=GOLD, bold=True, align=PP_ALIGN.CENTER)
    add_card(
        s,
        3.15,
        1.48,
        9.15,
        2.25,
        "Дополнительный фокус внимания",
        "Как терапевт, клиент и супервизионная группа соотносят происходящее с Богом: молитва, свобода, смирение, ответственность, различение духовного и психологического, избегание самовольного «знания воли Божией» за клиента.",
        GOLD_LIGHT,
        GOLD,
        body_size=14.2,
    )
    add_bullets(
        s,
        [
            "Не подмена профессионального анализа духовными объяснениями.",
            "Не давление на клиента религиозным языком или авторитетом.",
            "Различение: где моя тревога, где контрперенос, где нравственная граница, где молитвенное прошение о помощи.",
            "Позиция: «не мы, но через нас Господь помогает».",
        ],
        3.3,
        4.15,
        8.75,
        2.15,
        size=13.2,
    )
    add_footer(s, 11)

    s = setup_slide(prs)
    slides.append(s)
    add_title(s, "Пентаграмма Гингера", "Пять измерений личности как подсказка для супервизионного вопроса")
    add_pentagram(s, 4.0, 4.05, 2.0)
    add_card(
        s,
        7.15,
        1.72,
        5.35,
        4.65,
        "Как использовать в разборе",
        "В каждом случае можно проверять, какой уровень выпадает из внимания: телесный, эмоциональный, когнитивный, социальный или духовный. Для православной супервизии духовное измерение не растворяет остальные, а помогает собрать их в целостный взгляд на человека.",
        GREEN_LIGHT,
        GREEN,
        body_size=14.5,
    )
    add_footer(s, 12)

    s = setup_slide(prs)
    slides.append(s)
    add_title(s, "Семифокусная карта случая", "Практические вопросы для супервизии")
    questions = [
        ("1. Клиент", "Что происходит с клиентом? Каков запрос, страдание, ресурс, риск?"),
        ("2. Терапевт", "Какая гипотеза, стратегия, техника и профессиональная позиция?"),
        ("3. Процесс", "Что рождается между клиентом и терапевтом прямо сейчас?"),
        ("4. Супервизируемый", "Какие чувства, телесные реакции, страхи и переносы у терапевта?"),
        ("5. Отношения с супервизором", "Что повторяется в супервизии из клиентского процесса?"),
        ("6. Супервизор", "Какие реакции и гипотезы возникают у супервизора?"),
        ("7. Бог", "Где место молитвы, смирения, свободы и Промысла Божьего?"),
    ]
    y = 1.5
    for i, (title, body) in enumerate(questions):
        x = 0.85 if i < 4 else 6.95
        yy = y + (i % 4) * 1.25
        fill, accent = (GOLD_LIGHT, GOLD) if i == 6 else (BLUE_LIGHT, BLUE)
        add_card(s, x, yy, 5.45, 0.95, title, body, fill, accent, title_size=10.8, body_size=8.8)
    add_footer(s, 13)

    s = setup_slide(prs)
    slides.append(s)
    add_title(s, "Итог", "Супервизия как пространство соборного различения")
    add_card(
        s,
        1.05,
        1.75,
        11.25,
        3.1,
        "Семифокусная модель",
        "Сохраняет профессиональную строгость шестифокусной супервизии и добавляет православному психологу явный фокус на отношениях с Богом: молитвенное внимание, смирение, свободу, ответственность и любовь к клиенту без нарушения границ.",
        BLUE_LIGHT,
        BLUE,
        title_size=18,
        body_size=17,
    )
    add_text(s, "Практический вектор: качество помощи + безопасность клиента + духовная трезвость специалиста", 1.15, 5.25, 11.0, 0.48, size=17, color=GOLD, bold=True, align=PP_ALIGN.CENTER)
    add_footer(s, 14)

    prs.save(OUT)
    return OUT


if __name__ == "__main__":
    path = build()
    print(f"Created {path}")
